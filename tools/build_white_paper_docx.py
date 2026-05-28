from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "v14.1-white-paper-draft.md"
OUT = ROOT / "GSL-RC-White-Paper-v14.1-final-review.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text.strip())
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
            if row_idx == 0:
                set_cell_shading(cell, "D9EAF7")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_formatted_runs(paragraph, text):
    pos = 0
    for match in re.finditer(r"`([^`]+)`|\[([^\]]+)\]\(([^)]+)\)", text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        if match.group(1):
            run = paragraph.add_run(match.group(1))
            run.font.name = "Courier New"
            run.font.size = Pt(10.5)
        else:
            add_hyperlink(paragraph, match.group(2), match.group(3))
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        if run.font.size is None:
            run.font.size = Pt(12)


def add_paragraph(doc, text, style=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph(style=style)
    add_formatted_runs(p, text)
    p.alignment = align
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    return p


def is_table_start(lines, i):
    if i + 1 >= len(lines):
        return False
    return lines[i].strip().startswith("|") and re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", lines[i + 1])


def parse_table(lines, i):
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        line = lines[i].strip().strip("|")
        rows.append([cell.strip() for cell in line.split("|")])
        i += 1
    return [rows[0]] + rows[2:], i


def add_table(doc, rows):
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            set_cell_text(table.cell(r_idx, c_idx), text, bold=(r_idx == 0))
    style_table(table)
    doc.add_paragraph()


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color in [
        ("Title", 18, "17365D"),
        ("Subtitle", 14, "444444"),
        ("Heading 1", 15, "17365D"),
        ("Heading 2", 13, "1F4E79"),
        ("Heading 3", 12, "365F91"),
    ]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if "Heading" in name else 0)
        style.paragraph_format.space_after = Pt(6)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    configure_styles(doc)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines = []
    skip_next_caption = False
    i = 0

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue

        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Courier New"
                run.font.size = Pt(9.5)
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(8)
                in_code = False
                code_lines = []
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(raw)
            i += 1
            continue

        if is_table_start(lines, i):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue

        if line.startswith("# "):
            title = line[2:].strip()
            p = doc.add_paragraph(style="Title")
            add_formatted_runs(p, title)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            text = line[3:].strip()
            if text in {"Suggested Citation", "Document Production Note", "References and Data Sources to Preserve"}:
                doc.add_section(WD_SECTION.NEW_PAGE)
            add_paragraph(doc, text, style="Heading 1", align=WD_ALIGN_PARAGRAPH.LEFT)
        elif line.startswith("### "):
            add_paragraph(doc, line[4:].strip(), style="Heading 2", align=WD_ALIGN_PARAGRAPH.LEFT)
        elif line.startswith("#### "):
            add_paragraph(doc, line[5:].strip(), style="Heading 3", align=WD_ALIGN_PARAGRAPH.LEFT)
        elif line.startswith("!["):
            match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
            if match:
                image_path = ROOT / match.group(2)
                if image_path.exists():
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(str(image_path), width=Inches(5.9))
                else:
                    add_paragraph(doc, f"[Missing figure: {match.group(2)}]")
        elif line.startswith("Caption:"):
            p = add_paragraph(doc, line, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            for run in p.runs:
                run.italic = True
                run.font.size = Pt(10.5)
        elif re.match(r"^\d+\.\s+", line):
            text = re.sub(r"^\d+\.\s+", "", line)
            p = doc.add_paragraph(style="List Number")
            add_formatted_runs(p, text)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_formatted_runs(p, line[2:].strip())
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        else:
            add_paragraph(doc, line)
        i += 1

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.text = "Great Salt Lake RC White Paper v14.1 final review"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(100, 100, 100)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
