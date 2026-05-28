from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "GSL-RC-White-Paper-v14.1-final-review.docx"
PDF = ROOT / "GSL-RC-White-Paper-v14.1-final-review.pdf"


def main():
    doc = Document(DOCX)
    normal = doc.styles["Normal"]
    with ZipFile(DOCX) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    pdf = PdfReader(str(PDF))
    print(f"DOCX: {DOCX}")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")
    print(f"Embedded media files: {len(media)}")
    print(f"Normal style font: {normal.font.name}")
    print(f"Normal style size: {normal.font.size.pt if normal.font.size else 'unset'}")
    print(f"PDF: {PDF}")
    print(f"PDF pages: {len(pdf.pages)}")
    print(f"DOCX size bytes: {DOCX.stat().st_size}")
    print(f"PDF size bytes: {PDF.stat().st_size}")


if __name__ == "__main__":
    main()
